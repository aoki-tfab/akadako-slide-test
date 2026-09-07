# -*- coding: utf-8 -*-
"""pptx を原本にしたらどうなるかの検証。HTML版と同じ内容を PowerPoint の機能で作る。"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PX = 6350                      # 1920px 幅を 13.333in に対応させる
def E(v): return Emu(int(v * PX))
FONT = "BIZ UDPGothic"

prs = Presentation()
prs.slide_width  = E(1920)
prs.slide_height = E(1080)
blank = prs.slide_layouts[6]

def C(h): return RGBColor.from_string(h)

def box(sl, x, y, w, h, fill, line, lw=4, radius=True):
    shp = sl.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        E(x), E(y), E(w), E(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    shp.line.color.rgb = C(line); shp.line.width = Pt(lw)
    shp.text_frame.text = ""
    shp.shadow.inherit = False
    return shp

def txt(sl, x, y, w, h, s, size=40, bold=False, color="1A1A1A",
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = C(color); r.font.name = FONT
    return tb

def arrow_block(sl, x, y, w, h, color="3D3D3D"):
    shp = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, E(x), E(y), E(w), E(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = C(color)
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp

def arrow_line(sl, x, y, w, color="3D3D3D", lw=5):
    """線に矢じりを付ける（Canvaで手作りしているのと同じ作り方）"""
    from pptx.enum.shapes import MSO_CONNECTOR
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x), E(y), E(x + w), E(y))
    cn.line.color.rgb = C(color); cn.line.width = Pt(lw)
    # 矢じりは python-pptx に API が無いので XML を直接足す
    ln = cn.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    import lxml.etree as etree
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    return cn

def band(sl, title):
    b = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(100), E(70), E(1720), E(76))
    b.fill.solid(); b.fill.fore_color.rgb = C("F2F2F2")
    b.line.fill.background(); b.shadow.inherit = False
    bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(100), E(70), E(14), E(76))
    bar.fill.solid(); bar.fill.fore_color.rgb = C("888888")
    bar.line.fill.background(); bar.shadow.inherit = False
    txt(sl, 140, 70, 1600, 76, title, 44, True, anchor=MSO_ANCHOR.MIDDLE)

# ---- 1枚目: 4択（枠＋中央揃え＋上下中央） --------------------------------
s1 = prs.slides.add_slide(blank)
lab = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(100), E(70), E(220), E(64))
lab.fill.solid(); lab.fill.fore_color.rgb = C("C0392B")
lab.line.fill.background(); lab.shadow.inherit = False
txt(s1, 100, 70, 220, 64, "質問", 40, False, "FFFFFF", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
txt(s1, 100, 200, 1720, 200, "今日、スマホや端末の\nロックを解除しましたか？", 64, True)
for i, (n, t) in enumerate([("①","していない"),("②","1〜2回"),("③","3〜5回"),("④","6回以上")]):
    x = 100 + i*440
    box(s1, x, 470, 400, 130, "FFFFFF", "555555")
    txt(s1, x, 470, 400, 130, f"{n} {t}", 38, False, "1A1A1A", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
txt(s1, 100, 950, 1720, 60, "このあと「では、正確に何回か言える人？」と問う。", 28, False, "777777")

# ---- 2枚目: 矢印つきの図（ブロック矢印） ----------------------------------
s2 = prs.slides.add_slide(blank)
band(s2, "答えは ③ 銀行のコンピュータ（ブロック矢印）")
box(s2, 100, 210, 420, 140, "FAFAFA", "444444", 5)
txt(s2, 140, 230, 340, 100, "あなた\n暗証番号を入れる", 32, True)
arrow_block(s2, 560, 250, 200, 80)
box(s2, 800, 210, 440, 140, "FFE9C7", "D08000", 5)
txt(s2, 840, 230, 360, 100, "銀行のコンピュータ\n合っているか確かめる", 32, True, "C25E00")
a = arrow_block(s2, 1260, 250, 110, 80); a.rotation = 180
box(s2, 1380, 210, 440, 140, "EAF4FF", "3A7BD5", 5)
txt(s2, 1420, 230, 360, 100, "OK / NG が返る\nあの数秒は、返事待ち", 32, True, "14385F")
txt(s2, 100, 440, 1720, 260,
    "① なぜ、カードの中で確かめないのだろう？\n② 遠くのコンピュータに聞くとき、何が行き来している？\n③ その行き来を、他の人が見ることはできる？", 40)
box(s2, 100, 720, 1720, 130, "F7FAFD", "1F4E79", 5)
txt(s2, 148, 740, 1620, 100, "学習問題　私たちは、どうやって「本人だ」と確かめられているのだろうか。", 36, True, "1F4E79")

# ---- 3枚目: 線＋矢じり（Canvaの作り方に近い） ------------------------------
s3 = prs.slides.add_slide(blank)
band(s3, "同じ図を「線＋矢じり」で作った場合")
box(s3, 120, 320, 440, 160, "FAFAFA", "444444", 6)
txt(s3, 120, 320, 440, 160, "示す", 44, True, "1A1A1A", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
arrow_line(s3, 600, 400, 190)
box(s3, 820, 320, 440, 160, "FFE9C7", "D08000", 6)
txt(s3, 820, 320, 440, 160, "確かめる", 44, True, "C25E00", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
arrow_line(s3, 1300, 400, 190)
box(s3, 1520, 320, 280, 160, "EAF4FF", "3A7BD5", 6)
txt(s3, 1520, 320, 280, 160, "返る", 44, True, "14385F", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
txt(s3, 100, 620, 1720, 200, "行って、返ってくる。\n往復しないと、認証にならない。", 40)

prs.save("pptxtest.pptx")
print("pptxtest.pptx", prs.slide_width, prs.slide_height, len(prs.slides.__iter__.__self__._sldIdLst))
